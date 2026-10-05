"""Creation de ruches en nombre : creation rapide et import de tableur.

La creation rapide (gratuite) cree un rucher et ses ruches, ou ajoute des
ruches a un rucher existant, en une transaction. L'import (Premium) lit un
fichier CSV ou Excel, devine les colonnes, verifie chaque ligne et ne cree rien
tant qu'une ligne est en erreur. Le materiel est toujours du nouveau materiel :
le stock de l'Atelier ne bouge pas, et les reines se saisissent ensuite.
"""

from __future__ import annotations

import base64
import binascii
import csv
import difflib
import io
import re
import unicodedata
from uuid import UUID

from fastapi import HTTPException, status
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app import models, schemas
from app.materiel_flux import default_frame_count

FREE_RUCHERS_LIMIT = 2
KNOWN_FORMATS = ("dadant", "langstroth", "warre", "ruchette", "nucleus", "autre")
FORMAT_ALIASES = {
    "d": "dadant", "da": "dadant", "dad": "dadant", "db": "dadant", "blatt": "dadant",
    "l": "langstroth", "lang": "langstroth", "langs": "langstroth", "lg": "langstroth",
    "w": "warre", "war": "warre",
    "r": "ruchette", "rt": "ruchette", "ruch": "ruchette",
    "n": "nucleus", "nuc": "nucleus", "nuclei": "nucleus", "nucleis": "nucleus",
}
HEADER_HINTS: dict[str, tuple[str, ...]] = {
    "rucher": ("rucher", "emplacement", "site", "lieu", "apiary"),
    "identifiant": ("identifiant", "ruche", "numero", "num", "no", "nom", "id", "hive", "code"),
    "type_ruche": ("type", "usage", "role", "categorie"),
    "format": ("format", "modele", "model", "dimension"),
    "nombre_cadres": ("cadre", "cadres", "frames", "nb cadres"),
}
MAX_IMPORT_ROWS = 1000


def ensure_rucher_quota(db: Session, user: models.User, new_ruchers: int = 1) -> None:
    """Compte gratuit : au plus deux ruchers."""
    if user.is_premium or new_ruchers <= 0:
        return
    count = db.query(models.Rucher).filter(models.Rucher.user_id == user.id).count()
    if count + new_ruchers > FREE_RUCHERS_LIMIT:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Limite gratuite atteinte: passez premium pour creer plus de 2 ruchers",
        )


def _plain(value: str) -> str:
    text = unicodedata.normalize("NFKD", value or "")
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).strip().lower()


def normalize_format(raw: str | None) -> tuple[str | None, str | None]:
    """Format reconnu (ou None si vide) et message d'erreur eventuel."""
    text = _plain(raw or "")
    if not text:
        return None, None
    letters = re.sub(r"[^a-z]", "", text)
    if not letters:
        return None, f"format inconnu : {raw}"
    if letters in KNOWN_FORMATS:
        return letters, None
    if letters in FORMAT_ALIASES:
        return FORMAT_ALIASES[letters], None
    for known in KNOWN_FORMATS:
        if len(letters) >= 3 and (known.startswith(letters) or letters.startswith(known)):
            return known, None
    close = difflib.get_close_matches(letters, KNOWN_FORMATS, n=1, cutoff=0.75)
    if close:
        return close[0], None
    return None, f"format inconnu : {raw}"


def _existing_identifiers(db: Session, user_id: UUID) -> set[str]:
    rows = (
        db.query(func.lower(func.trim(models.Ruche.identifiant_personnalise)))
        .filter(models.Ruche.user_id == user_id, models.Ruche.archived_at.is_(None))
        .all()
    )
    return {row[0] for row in rows}


def hive_types(db: Session, user_id: UUID) -> dict[str, UUID]:
    """Types de ruche du compte (systeme et personnels), par libelle normalise."""
    rows = (
        db.query(models.RefTypeRuche)
        .filter(or_(models.RefTypeRuche.is_system.is_(True), models.RefTypeRuche.user_id == user_id))
        .order_by(models.RefTypeRuche.is_system.asc())
        .all()
    )
    return {_plain(row.libelle): row.id for row in rows}


def validate_hive_type(db: Session, user_id: UUID, type_id: UUID | None) -> None:
    if type_id is not None and type_id not in hive_types(db, user_id).values():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Type de ruche introuvable")


def match_hive_type(types: dict[str, UUID], raw: str | None) -> tuple[UUID | None, str | None]:
    """Type reconnu par son libelle (accents, casse, pluriel, faute proche)."""
    text = _plain(raw or "")
    if not text:
        return None, None
    if text in types:
        return types[text], None
    singular = text[:-1] if text.endswith("s") else text
    for label, type_id in types.items():
        if singular == label or label.startswith(singular) and len(singular) >= 4:
            return type_id, None
    close = difflib.get_close_matches(text, list(types), n=1, cutoff=0.8)
    if close:
        return types[close[0]], None
    return None, f"type inconnu : {raw}"


def _active_status_id(db: Session, user_id: UUID) -> UUID | None:
    from app.routers.ruches import _default_active_status_id

    return _default_active_status_id(db, user_id)


def create_hives(db: Session, user: models.User, rucher: models.Rucher, hives: list[schemas.RucheRapide]) -> int:
    """Cree les ruches dans le rucher ; refuse tout si un identifiant est pris."""
    existing = _existing_identifiers(db, user.id)
    seen: set[str] = set()
    duplicates: list[str] = []
    for hive in hives:
        key = hive.identifiant_personnalise.strip().lower()
        if key in existing or key in seen:
            duplicates.append(hive.identifiant_personnalise.strip())
        seen.add(key)
    if duplicates:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"Identifiant(s) deja utilise(s) : {', '.join(duplicates[:10])}",
        )
    status_id = _active_status_id(db, user.id)
    for type_id in {hive.ref_type_ruche_id for hive in hives if hive.ref_type_ruche_id}:
        validate_hive_type(db, user.id, type_id)
    for hive in hives:
        format_ruche = hive.format_ruche or None
        frames = hive.nombre_cadres
        if frames is None and format_ruche:
            frames = default_frame_count(format_ruche)
        db.add(models.Ruche(
            user_id=user.id,
            rucher_id=rucher.id,
            is_at_atelier=False,
            identifiant_personnalise=hive.identifiant_personnalise.strip(),
            format_ruche=format_ruche,
            nombre_cadres=frames,
            ref_statut_ruche_id=status_id,
            ref_type_ruche_id=hive.ref_type_ruche_id,
            has_corps=hive.has_corps,
            has_toit=hive.has_toit,
            has_plancher=hive.has_plancher,
            has_grille_a_reine=hive.has_grille_a_reine,
            has_nourrisseur=hive.has_nourrisseur,
            has_hausse=hive.has_hausse,
            has_partition=hive.has_partition,
        ))
    db.flush()
    return len(hives)


# --- Import de tableur ------------------------------------------------------------

def _decode_file(nom_fichier: str, contenu_base64: str) -> list[list[str]]:
    try:
        raw = base64.b64decode(contenu_base64, validate=True)
    except (binascii.Error, ValueError):
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Fichier illisible")
    name = nom_fichier.lower()
    if name.endswith(".xls"):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Ancien format Excel (.xls) : enregistrer le fichier en .xlsx ou .csv",
        )
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - fichier corrompu ou non Excel
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Fichier Excel illisible")
        sheet = workbook.worksheets[0]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else _cell_text(cell) for cell in row])
            if len(rows) > MAX_IMPORT_ROWS + 1:
                break
        workbook.close()
        return rows
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Encodage du fichier non reconnu")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _cell_text(cell) -> str:
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


def guess_mapping(headers: list[str]) -> dict[str, str | None]:
    """Colonne la plus probable pour chaque champ, une colonne au plus par champ."""
    mapping: dict[str, str | None] = {field: None for field in HEADER_HINTS}
    used: set[str] = set()
    plain_headers = [(header, _plain(header)) for header in headers if header.strip()]
    # Le rucher d'abord : "Nom du rucher" ne doit pas devenir l'identifiant.
    for field in ("rucher", "nombre_cadres", "format", "type_ruche", "identifiant"):
        for header, plain in plain_headers:
            if header in used:
                continue
            words = set(re.findall(r"[a-z]+", plain))
            if any(hint == plain or hint in words or (len(hint) > 3 and hint in plain) for hint in HEADER_HINTS[field]):
                if field == "identifiant" and "rucher" in plain:
                    continue
                mapping[field] = header
                used.add(header)
                break
    return mapping


def analyse_file(payload: schemas.ImportAnalyseRequest) -> tuple[list[str], dict[str, str | None], list[schemas.LigneImport]]:
    rows = [row for row in _decode_file(payload.nom_fichier, payload.contenu_base64) if any(cell.strip() for cell in row)]
    if not rows:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Fichier vide")
    headers = [cell.strip() for cell in rows[0]]
    if len(rows) - 1 > MAX_IMPORT_ROWS:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Fichier trop long : {MAX_IMPORT_ROWS} ruches au plus par import",
        )
    mapping = guess_mapping(headers)
    # Une colonne "Type" contient souvent le format (Dadant...) : si ses valeurs
    # sont pour moitie au moins des formats et qu'aucune colonne format n'est trouvee, elle
    # devient la colonne format.
    type_column = mapping.get("type_ruche")
    if type_column and not mapping.get("format"):
        position = headers.index(type_column)
        values = [row[position] for row in rows[1:] if position < len(row) and row[position].strip()]
        formats = sum(1 for item in values if normalize_format(item)[0])
        if values and formats * 2 >= len(values):
            mapping["format"], mapping["type_ruche"] = type_column, None
    if payload.correspondance:
        mapping.update({field: (column or None) for field, column in payload.correspondance.items()})
    index = {field: (headers.index(column) if column in headers else None) for field, column in mapping.items()}

    def value(row: list[str], field: str) -> str:
        position = index[field]
        return row[position].strip() if position is not None and position < len(row) else ""

    lines = [
        schemas.LigneImport(
            rucher=value(row, "rucher")[:150],
            identifiant=value(row, "identifiant")[:100],
            type_ruche=value(row, "type_ruche")[:100] or None,
            format_ruche=value(row, "format")[:30] or None,
            nombre_cadres=value(row, "nombre_cadres")[:10] or None,
        )
        for row in rows[1:]
    ]
    return headers, mapping, lines


def verify_lines(db: Session, user: models.User, lines: list[schemas.LigneImport]) -> schemas.ImportVerificationResponse:
    existing_ids = _existing_identifiers(db, user.id)
    existing_ruchers = {
        _plain(rucher.nom)
        for rucher in db.query(models.Rucher).filter(models.Rucher.user_id == user.id).all()
    }
    types = hive_types(db, user.id)
    seen: dict[str, int] = {}
    new_ruchers: dict[str, str] = {}
    checked: list[schemas.LigneImportVerifiee] = []
    for position, line in enumerate(lines, start=2):
        errors: list[str] = []
        rucher = re.sub(r"\s+", " ", line.rucher).strip()
        identifiant = line.identifiant.strip()
        if not rucher:
            errors.append("rucher manquant")
        if not identifiant:
            errors.append("identifiant manquant")
        else:
            key = identifiant.lower()
            if key in existing_ids:
                errors.append(f"identifiant {identifiant} deja utilise dans Beevarium")
            elif key in seen:
                errors.append(f"identifiant {identifiant} en double (ligne {seen[key]})")
            else:
                seen[key] = position
        format_ruche, format_error = normalize_format(line.format_ruche)
        if format_error:
            errors.append(format_error)
        _, type_error = match_hive_type(types, line.type_ruche)
        if type_error:
            errors.append(type_error)
        frames_text = (line.nombre_cadres or "").strip().replace(",", ".")
        frames_value = ""
        if frames_text:
            try:
                number = float(frames_text)
                if not number.is_integer() or not 0 <= number <= 100:
                    raise ValueError
                frames_value = str(int(number))
            except ValueError:
                errors.append(f"nombre de cadres invalide : {line.nombre_cadres}")
                frames_value = line.nombre_cadres or ""
        if rucher and _plain(rucher) not in existing_ruchers:
            new_ruchers.setdefault(_plain(rucher), rucher)
        checked.append(schemas.LigneImportVerifiee(
            numero=position,
            rucher=rucher,
            identifiant=identifiant,
            type_ruche=line.type_ruche,
            format_ruche=format_ruche or line.format_ruche,
            nombre_cadres=frames_value or None,
            erreurs=errors,
            rucher_existant=bool(rucher) and _plain(rucher) in existing_ruchers,
        ))
    return schemas.ImportVerificationResponse(
        lignes=checked,
        nb_erreurs=sum(1 for line in checked if line.erreurs),
        ruchers_a_creer=list(new_ruchers.values()),
    )


def import_lines(db: Session, user: models.User, lines: list[schemas.LigneImport]) -> schemas.ImportResultat:
    verification = verify_lines(db, user, lines)
    if verification.nb_erreurs:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"{verification.nb_erreurs} ligne(s) en erreur : corriger avant d'importer",
        )
    ruchers = {
        _plain(rucher.nom): rucher
        for rucher in db.query(models.Rucher).filter(models.Rucher.user_id == user.id).all()
    }
    ensure_rucher_quota(db, user, len(verification.ruchers_a_creer))
    created_ruchers = 0
    for name in verification.ruchers_a_creer:
        rucher = models.Rucher(user_id=user.id, nom=name, statut_activite="actif", statut_peuplement="peuple")
        db.add(rucher)
        db.flush()
        ruchers[_plain(name)] = rucher
        created_ruchers += 1
    types = hive_types(db, user.id)
    groups: dict[str, list[schemas.RucheRapide]] = {}
    for line in verification.lignes:
        groups.setdefault(_plain(line.rucher), []).append(schemas.RucheRapide(
            identifiant_personnalise=line.identifiant,
            ref_type_ruche_id=match_hive_type(types, line.type_ruche)[0],
            format_ruche=line.format_ruche,
            nombre_cadres=int(line.nombre_cadres) if line.nombre_cadres else None,
        ))
    created = sum(create_hives(db, user, ruchers[key], hives) for key, hives in groups.items())
    return schemas.ImportResultat(ruches_creees=created, ruchers_crees=created_ruchers)
